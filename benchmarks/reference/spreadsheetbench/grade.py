"""Grade a driver.py run with SpreadsheetBench's own comparison, evaluation.compare_workbooks
(evaluation/evaluation.py of the upstream checkout, imported unmodified), two ways:

 raw : the workbook as the model's code wrote it, read by openpyxl with data_only=True, against
       the shipped golden file. This is upstream's grading as coded: a formula that openpyxl wrote
       has no cached value, so it reads as empty.
 lo  : the workbook after a LibreOffice recalculation (the conversion upstream's
       evaluation/open_spreadsheet.py runs: soffice --headless --calc --convert-to
       "xlsx:Calc MS Excel 2007 XML"), against the golden file put through the same conversion.
       lo_rawgold: the recalculated workbook against the shipped golden file.

Upstream's evaluation.py main cannot score verified_400 as shipped (it reads the task's input
file instead of the output, and expects earlier file names), so only compare_workbooks is used.

usage: python3 grade.py TAG [IDS_FILE]
Reads $SBREF_OUT/TAG/outputs (default ~/.cache/dvergr-bench/sbref/TAG), writes
$SBREF_OUT/TAG/grade.json ({id: {exists, raw, lo, lo_rawgold, uncached, ...}}).
Environment: SPREADSHEETBENCH_ROOT, SPREADSHEETBENCH_REPO, SBREF_OUT as for driver.py;
  LIBREOFFICE   the LibreOffice binary (default libreoffice)
  SBREF_LO_STAGE  scratch dir for the conversion (default ~/snap/libreoffice/common/sbref-stage
                  when LibreOffice is a snap, which can only read and write there, else
                  $SBREF_OUT/lo-stage)
"""
import contextlib, glob, importlib.util, io, json, os, shutil, subprocess, sys

import openpyxl

HERE = os.path.dirname(os.path.abspath(__file__))
env = lambda k, d: os.environ.get(k) or d
REPO = os.path.expanduser(env('SPREADSHEETBENCH_REPO', '~/.cache/dvergr-bench/spreadsheetbench/repo'))
D = os.path.expanduser(env('SPREADSHEETBENCH_ROOT',
                           '~/.cache/dvergr-bench/spreadsheetbench/data/spreadsheetbench_verified_400'))
OUT = os.path.expanduser(env('SBREF_OUT', '~/.cache/dvergr-bench/sbref'))
LO = env('LIBREOFFICE', 'libreoffice')
_lo_path = os.path.realpath(shutil.which(LO) or LO)
LOS = os.path.expanduser(env('SBREF_LO_STAGE', '~/snap/libreoffice/common/sbref-stage'
                             if '/snap/' in _lo_path else f'{OUT}/lo-stage'))
spec = importlib.util.spec_from_file_location('sb_upstream_eval', f'{REPO}/evaluation/evaluation.py')
ev = importlib.util.module_from_spec(spec)
spec.loader.exec_module(ev)


def gold(tid):
    f = f'{D}/spreadsheet/{tid}/1_{tid}_golden.xlsx'
    return f if os.path.exists(f) else f'{D}/spreadsheet/{tid}/golden.xlsx'


def cmp(gt, proc, t):
    with contextlib.redirect_stdout(io.StringIO()):
        try:
            ok, _ = ev.compare_workbooks(gt, proc, t['instruction_type'], t['answer_position'])
        except Exception as e:
            return False, 'EXC ' + repr(e)[:150]
        if ok:
            return True, ''
        msgs = []
        try:
            wg = openpyxl.load_workbook(gt, data_only=True); wp = openpyxl.load_workbook(proc, data_only=True)
            for scr in t['answer_position'].split(','):
                if '!' in scr:
                    sn, cr = scr.split('!')
                else:
                    sn, cr = wg.sheetnames[0], scr
                o, m = ev.cell_level_compare(wg, wp, sn.strip("'"), cr.strip("'"))
                if not o:
                    msgs.append(' '.join(m.split()))
        except Exception as e:
            msgs.append('EXC ' + repr(e)[:150])
    return False, ' | '.join(msgs)[:300]


def uncached_formulas(proc, t):
    """Formula cells in the answer range that carry no cached value (openpyxl-written formulas)."""
    try:
        wf = openpyxl.load_workbook(proc); wv = openpyxl.load_workbook(proc, data_only=True)
    except Exception:
        return None
    n = 0
    for scr in t['answer_position'].split(','):
        if '!' in scr:
            sn, cr = scr.split('!')
        else:
            sn, cr = wf.sheetnames[0], scr
        sn, cr = sn.strip("'"), cr.strip("'")
        if sn not in wf.sheetnames:
            continue
        for c in ev.generate_cell_names(cr):
            f = wf[sn][c].value
            if isinstance(f, str) and f.startswith('=') or type(f).__name__ in ('ArrayFormula', 'DataTableFormula'):
                if wv[sn][c].value is None:
                    n += 1
    return n


def recalc(src_dir, dst_dir):
    """Convert every workbook of src_dir into dst_dir; returns the names LibreOffice did not convert."""
    files = sorted(glob.glob(f'{src_dir}/*.xlsx'))
    if files:
        p = subprocess.run([LO, f'-env:UserInstallation=file://{LOS}/profile', '--headless', '--calc',
                            '--convert-to', 'xlsx:Calc MS Excel 2007 XML', '--outdir', dst_dir] + files,
                           capture_output=True, text=True, timeout=3600)
        if p.returncode != 0:
            print(f'LibreOffice exited with {p.returncode}:', p.stdout[-2000:], p.stderr[-2000:], file=sys.stderr)
    return [os.path.basename(f) for f in files if not os.path.exists(f'{dst_dir}/{os.path.basename(f)}')]


def main():
    if len(sys.argv) < 2:
        sys.exit(__doc__)
    tag = sys.argv[1]
    ids = open(sys.argv[2] if len(sys.argv) > 2 else os.path.join(HERE, 'ids.txt')).read().split()
    ds = {str(d['id']): d for d in json.load(open(f'{D}/dataset.json'))}
    run_dir = f'{OUT}/{tag}'
    stage = f'{LOS}/{tag}'
    shutil.rmtree(stage, ignore_errors=True)
    for d in ('in', 'lo', 'gold-in', 'gold'):
        os.makedirs(f'{stage}/{d}')
    outs = {t: f'{run_dir}/outputs/{t}.xlsx' for t in ids}
    for t, f in outs.items():
        if os.path.exists(f):
            shutil.copy(f, f'{stage}/in/{t}.xlsx')
        shutil.copy(gold(t), f'{stage}/gold-in/{t}.xlsx')
    failed = recalc(f'{stage}/in', f'{stage}/lo') + recalc(f'{stage}/gold-in', f'{stage}/gold')
    if failed:
        # a missing conversion would be scored as a wrong answer: stop instead
        sys.exit(f'LibreOffice did not convert {len(failed)} workbooks ({failed[:10]}); no grades written')
    res = {}
    for t in ids:
        d = ds[t]; out = outs[t]; lo = f'{stage}/lo/{t}.xlsx'; lo_gold = f'{stage}/gold/{t}.xlsx'
        r = {'exists': os.path.exists(out)}
        r['raw'], r['raw_msg'] = cmp(gold(t), out, d) if r['exists'] else (False, 'no output')
        r['lo'], r['lo_msg'] = cmp(lo_gold, lo, d) if os.path.exists(lo) else (False, 'no LO output')
        r['lo_rawgold'], _ = cmp(gold(t), lo, d) if os.path.exists(lo) else (False, '')
        try:
            r['uncached'] = uncached_formulas(out, d) if r['exists'] else None
        except Exception:
            r['uncached'] = None
        res[t] = r
    json.dump(res, open(f'{run_dir}/grade.json', 'w'), indent=1)
    n = sum(1 for r in res.values())
    print(tag, 'tasks', n, 'with output', sum(r['exists'] for r in res.values()),
          '| raw', sum(r['raw'] for r in res.values()), '| lo', sum(r['lo'] for r in res.values()),
          '| lo_rawgold', sum(r['lo_rawgold'] for r in res.values()),
          '| tasks with uncached formulas', sum(1 for r in res.values() if r['uncached']))


if __name__ == '__main__':
    main()
